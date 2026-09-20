"""Agent 3 — Document Parser.

Deterministic JSON/CSV parsers; unstructured formats go through Claude Haiku
with evidence snippets. Public entrypoint: parse_response(path, vendor_id, rfx).
"""
from __future__ import annotations

import base64
import csv
import email
import json
import os
import re
from email import policy
from pathlib import Path
from typing import Any, Optional, Union

from shared_models import ExtractedQuote, RFx

_HAIKU_FALLBACKS = [
    "claude-3-haiku-20240307",
    "claude-haiku-4-5-20251001",
    "claude-3-5-haiku-20241022",
]


def _resolve_haiku_model() -> str:
    preferred = os.environ.get("ANTHROPIC_HAIKU_MODEL", "").strip()
    for c in [preferred, *_HAIKU_FALLBACKS]:
        if c:
            return c
    return "claude-3-haiku-20240307"


HAIKU_MODEL = _resolve_haiku_model()

JSON_CSV_EXTS = {".json", ".csv"}
TEXT_EXTS = {".txt", ".md", ".eml", ".msg"}
DOC_EXTS = {".docx", ".pdf"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
SUPPORTED = JSON_CSV_EXTS | TEXT_EXTS | DOC_EXTS | IMAGE_EXTS

SYSTEM_PROMPT = """You are a procurement document parser for vendor RFx replies.
Extract ONLY facts present in the vendor document. Never invent prices, line items, or answers.
If the vendor says "same as last year", "freight extra", ambiguous UOMs (per kg without weight, per 100 pcs, etc.), flag those clearly in notes and on the affected lines.
Return JSON ONLY matching this schema:
{
  "vendor_name": "string or empty",
  "currency": "INR|USD|EUR|...",
  "lines": [
    {
      "line_id": "string (RFx line id or number if stated)",
      "description": "string",
      "unit_price": number or null,
      "uom": "string",
      "notes": "string",
      "currency": "string"
    }
  ],
  "questionnaire_answers": [{"question": "string", "answer": "string"}],
  "notes": "commercial caveats, freight, same-as-last-year, missing items",
  "confidence": 0.0-1.0,
  "raw_evidence": [{"snippet": "verbatim quote", "location": "optional"}],
  "extraction_notes": "ambiguities and what was omitted"
}
Rules:
- Include a line only when a concrete price or explicit refusal appears in the text.
- Prefer matching descriptions to the RFx line list when provided; leave line_id empty if unsure.
- Do not invent questionnaire answers.
- JSON only — no markdown fences.
"""


def _guess_vendor_id(path: Path, vendor_id: Optional[str] = None, vendor_hint: str = "") -> str:
    if vendor_id:
        return str(vendor_id).strip()
    if vendor_hint:
        m = re.search(r"(V\d+)", vendor_hint, re.I)
        if m:
            return m.group(1).upper()
        return vendor_hint.strip()[:32]
    m = re.search(r"(V\d+)", path.stem, re.I)
    return m.group(1).upper() if m else path.stem[:32]


def _rfx_context(rfx: Any) -> dict[str, Any]:
    if rfx is None:
        return {"rfx_id": "", "currency": "INR", "line_items": [], "questionnaire": []}
    if isinstance(rfx, RFx):
        data = rfx.model_dump()
    elif isinstance(rfx, dict):
        data = rfx
    elif hasattr(rfx, "model_dump"):
        data = rfx.model_dump()
    else:
        data = {}
    lines = data.get("line_items") or data.get("lines") or []
    slim_lines = []
    for li in lines:
        if isinstance(li, dict):
            slim_lines.append(
                {
                    "line_id": str(li.get("line_id") or li.get("id") or li.get("line_no") or ""),
                    "description": str(li.get("description") or li.get("name") or ""),
                    "uom": str(li.get("uom") or "piece"),
                    "qty": li.get("qty"),
                }
            )
        else:
            slim_lines.append(
                {
                    "line_id": str(getattr(li, "line_id", "")),
                    "description": str(getattr(li, "description", "")),
                }
            )
    qs = data.get("questionnaire") or []
    slim_q = []
    for q in qs:
        if isinstance(q, dict):
            slim_q.append(
                {
                    "id": q.get("id", ""),
                    "question": q.get("question") or q.get("text") or "",
                    "knockout": bool(q.get("knockout", False)),
                }
            )
        else:
            slim_q.append(
                {
                    "id": getattr(q, "id", ""),
                    "question": getattr(q, "question", ""),
                    "knockout": bool(getattr(q, "knockout", False)),
                }
            )
    return {
        "rfx_id": str(data.get("rfx_id") or data.get("id") or ""),
        "currency": str(data.get("currency") or "INR"),
        "line_items": slim_lines,
        "questionnaire": slim_q,
    }


def _as_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "").replace("₹", "").replace("$", "").replace("€", "")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return float(m.group(0)) if m else None


def _normalize_line(row: dict[str, Any], default_currency: str = "INR") -> dict[str, Any]:
    line_id = row.get("line_id")
    if line_id is None:
        line_id = row.get("Line#") or row.get("line") or row.get("Line") or row.get("id") or ""
    line_id = str(line_id).strip()

    description = (
        row.get("description")
        or row.get("Item Description")
        or row.get("item")
        or row.get("name")
        or row.get("sku_name")
        or ""
    )

    price = None
    for key in (
        "unit_price",
        "unit_price_inr",
        "Unit Price (INR)",
        "Unit Price",
        "price",
        "rate",
        "amount",
        "raw_price",
    ):
        if key in row and row[key] not in (None, ""):
            price = _as_float(row[key])
            break

    uom = row.get("uom") or row.get("UOM") or row.get("unit") or ""
    notes = row.get("notes") or row.get("Remarks") or row.get("remarks") or ""
    currency = row.get("currency") or default_currency

    out: dict[str, Any] = {
        "line_id": line_id,
        "description": str(description).strip(),
        "unit_price": price,
        "uom": str(uom).strip(),
        "notes": str(notes).strip(),
        "currency": str(currency).strip() or default_currency,
    }
    if "unit_price_inr" in row and price is not None:
        out["unit_price_inr"] = price
    return out


def _normalize_answers(raw: Any) -> list[dict[str, Any]]:
    """Normalize raw questionnaire payload into {id,question,answer,answered} dicts."""
    answers: list[dict[str, Any]] = []
    if isinstance(raw, dict):
        for q, a in raw.items():
            ans = "" if a is None else str(a)
            answers.append(
                {
                    "id": "",
                    "question": str(q),
                    "answer": ans,
                    "answered": bool(str(ans).strip()),
                }
            )
    elif isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                qid = str(item.get("id") or "")
                q = item.get("question") or item.get("q") or item.get("text") or ""
                if not q and qid and not item.get("answer") and not item.get("a"):
                    # id-only key misuse — treat id as question label
                    q = qid
                    qid = qid if re.match(r"^Q\d+", str(qid), re.I) else ""
                a = item.get("answer") or item.get("a") or item.get("response") or ""
                ans = "" if a is None else str(a)
                # If id looks like Qn and question empty, keep id
                if not qid and re.match(r"^Q\d+", str(item.get("id") or ""), re.I):
                    qid = str(item.get("id"))
                answers.append(
                    {
                        "id": qid,
                        "question": str(q),
                        "answer": ans,
                        "answered": bool(ans.strip()),
                    }
                )
            else:
                answers.append({"id": "", "question": "", "answer": str(item), "answered": bool(str(item).strip())})
    return answers


def _norm_qtext(s: str) -> str:
    s = (s or "").lower().strip()
    s = re.sub(r"[^a-z0-9\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _quality_gate_answers(
    answers: list[dict[str, Any]],
    rfx_ctx: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Map extracted answers onto RFx questionnaire ids for knockout gating.

    Never invents answers. Unanswered knockouts (when RFx provided) appear with
    answered=False and empty answer.
    """
    rfx_ctx = rfx_ctx or {}
    rfx_qs = list(rfx_ctx.get("questionnaire") or [])

    def _tokens(s: str) -> set[str]:
        # Drop ultra-common procurement filler words to avoid false matches
        stop = {
            "the", "and", "for", "you", "your", "are", "is", "this", "that",
            "with", "from", "have", "will", "can", "into", "within", "over",
            "last", "days", "day", "per", "any", "all", "please", "provide",
        }
        return {t for t in _norm_qtext(s).split() if len(t) > 2 and t not in stop}

    def _match_rfx(qid: str, qtext: str, used: set[str]):
        if not rfx_qs:
            return None
        qn = _norm_qtext(qtext)
        # 1) explicit id
        if qid:
            for rq in rfx_qs:
                rq_id = str(rq.get("id") or "")
                if rq_id and rq_id.upper() == qid.upper() and rq_id.upper() not in used:
                    return rq
        # 2) strong text containment / equality
        if qn:
            for rq in rfx_qs:
                rq_id = str(rq.get("id") or "")
                if rq_id.upper() in used:
                    continue
                rn = _norm_qtext(str(rq.get("question") or ""))
                if not rn:
                    continue
                if qn == rn or qn in rn or rn in qn:
                    return rq
        # 3) conservative token overlap (need >=3 shared content tokens AND >=50% of smaller side)
        q_tokens = _tokens(qtext)
        if len(q_tokens) >= 3:
            best = None
            best_score = 0
            for rq in rfx_qs:
                rq_id = str(rq.get("id") or "")
                if rq_id.upper() in used:
                    continue
                r_tokens = _tokens(str(rq.get("question") or ""))
                if len(r_tokens) < 3:
                    continue
                score = len(q_tokens & r_tokens)
                smaller = min(len(q_tokens), len(r_tokens))
                if score >= 3 and score >= ceil_half(smaller) and score > best_score:
                    best_score = score
                    best = rq
            return best
        return None

    def ceil_half(n: int) -> int:
        return (n + 1) // 2

    shaped: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    for raw in answers:
        if not isinstance(raw, dict):
            continue
        qid = str(raw.get("id") or "")
        qtext = str(raw.get("question") or "")
        ans = "" if raw.get("answer") is None else str(raw.get("answer"))
        best = _match_rfx(qid, qtext, used_ids)
        if best is not None:
            qid = str(best.get("id") or qid)
            qtext = str(best.get("question") or qtext)
            if qid:
                used_ids.add(qid.upper())
        shaped.append(
            {
                "id": qid,
                "question": qtext,
                "answer": ans,
                "answered": bool(ans.strip()),
            }
        )

    # Ensure knockout questions appear even when vendor omitted them (empty, not invented)
    if rfx_qs:
        seen_ids = {str(s.get("id") or "").upper() for s in shaped if s.get("id")}
        seen_q = {_norm_qtext(str(s.get("question") or "")) for s in shaped}
        for rq in rfx_qs:
            if not bool(rq.get("knockout")):
                continue
            rq_id = str(rq.get("id") or "")
            rq_text = str(rq.get("question") or "")
            if rq_id and rq_id.upper() in seen_ids:
                continue
            if rq_text and _norm_qtext(rq_text) in seen_q:
                continue
            shaped.append(
                {
                    "id": rq_id,
                    "question": rq_text,
                    "answer": "",
                    "answered": False,
                }
            )
    return shaped



def _meta_evidence(
    path: Path,
    *,
    parse_method: str,
    vendor_name: str = "",
    currency: str = "INR",
    extraction_notes: str = "",
    extra: Optional[list[dict[str, Any]]] = None,
) -> list[dict[str, Any]]:
    meta = {
        "kind": "meta",
        "parse_method": parse_method,
        "vendor_name": vendor_name,
        "source_file": path.name,
        "currency": currency,
        "extraction_notes": extraction_notes,
    }
    out = [meta]
    if extra:
        out.extend(extra)
    return out


def _quote(
    *,
    vendor_id: str,
    source_format: str,
    lines: list[dict[str, Any]],
    questionnaire_answers: list[dict[str, Any]],
    notes: str,
    confidence: float,
    raw_evidence: list[dict[str, Any]],
) -> ExtractedQuote:
    return ExtractedQuote(
        vendor_id=vendor_id,
        source_format=source_format,
        lines=lines,
        questionnaire_answers=questionnaire_answers,
        notes=notes,
        confidence=confidence,
        raw_evidence=raw_evidence,
    )


def _parse_json(path: Path, vendor_id: str) -> ExtractedQuote:
    data = json.loads(path.read_text(encoding="utf-8"))
    currency = str(data.get("currency") or "INR")
    vendor_name = str(data.get("vendor") or data.get("vendor_name") or "")
    raw_lines = data.get("line_items") or data.get("lines") or data.get("items") or []
    lines = [_normalize_line(dict(row), currency) for row in raw_lines if isinstance(row, dict)]
    answers = _normalize_answers(
        data.get("questionnaire")
        or data.get("questionnaire_answers")
        or data.get("answers")
        or {}
    )
    notes = str(data.get("notes") or "")
    if data.get("validity_days"):
        notes = (notes + f" Validity: {data['validity_days']} days.").strip()
    evidence = _meta_evidence(
        path,
        parse_method="deterministic_json",
        vendor_name=vendor_name,
        currency=currency,
        extra=[{"snippet": path.name, "location": "file", "kind": "json"}],
    )
    return _quote(
        vendor_id=str(data.get("vendor_id") or vendor_id),
        source_format="json",
        lines=lines,
        questionnaire_answers=answers,
        notes=notes,
        confidence=0.95,
        raw_evidence=evidence,
    )


def _parse_csv(path: Path, vendor_id: str) -> ExtractedQuote:
    text = path.read_text(encoding="utf-8-sig")
    lines_out: list[dict[str, Any]] = []
    answers: list[dict[str, Any]] = []
    evidence_bits: list[dict[str, Any]] = []
    section = "lines"
    header: Optional[list[str]] = None
    currency = "INR"
    vendor_name = ""

    for lineno, raw in enumerate(text.splitlines(), start=1):
        row = raw.strip()
        if not row:
            continue
        upper = row.upper()
        if "QUESTIONNAIRE" in upper:
            section = "questionnaire"
            header = None
            evidence_bits.append({"snippet": row, "location": f"line {lineno}"})
            continue
        if section == "questionnaire":
            if "," in row:
                parts = next(csv.reader([row]))
                if len(parts) >= 2:
                    answers.append({"question": parts[0].strip(), "answer": ",".join(parts[1:]).strip()})
                else:
                    answers.append({"question": parts[0].strip(), "answer": ""})
            elif ":" in row:
                q, a = row.split(":", 1)
                answers.append({"question": q.strip(), "answer": a.strip()})
            else:
                answers.append({"question": row, "answer": ""})
            continue

        if header is None:
            header = [h.strip() for h in next(csv.reader([row]))]
            continue
        cells = next(csv.reader([row]))
        while len(cells) < len(header):
            cells.append("")
        mapped = {header[i]: cells[i].strip() for i in range(len(header))}
        if any(str(v).lower() in {"line#", "line_id", "item description"} for v in mapped.values()):
            continue
        norm = _normalize_line(mapped, currency)
        if not norm["description"] and norm["unit_price"] is None:
            continue
        lines_out.append(norm)
        if norm["unit_price"] is not None:
            evidence_bits.append({"snippet": raw[:200], "location": f"line {lineno}"})

    stem = path.stem
    m = re.match(r"V\d+[_-]?(.*)$", stem, re.I)
    if m and m.group(1):
        vendor_name = m.group(1).replace("_", " ").replace("-", " ").strip()
        vendor_name = re.sub(r"\s*response\s*$", "", vendor_name, flags=re.I).strip()

    notes_parts = []
    messy_uoms = [ln for ln in lines_out if re.search(r"100|kg|bundle|lot", ln.get("uom", ""), re.I)]
    if messy_uoms:
        notes_parts.append(f"{len(messy_uoms)} line(s) use non-piece UOMs (per 100 / kg / etc.).")

    evidence = _meta_evidence(
        path,
        parse_method="deterministic_csv",
        vendor_name=vendor_name,
        currency=currency,
        extra=evidence_bits[:40],
    )
    return _quote(
        vendor_id=vendor_id,
        source_format="csv",
        lines=lines_out,
        questionnaire_answers=answers,
        notes=" ".join(notes_parts),
        confidence=0.9 if lines_out else 0.5,
        raw_evidence=evidence,
    )


def _extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_eml(path: Path) -> str:
    msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    parts: list[str] = []
    for header in ("From", "To", "Subject", "Date"):
        val = msg.get(header)
        if val:
            parts.append(f"{header}: {val}")
    parts.append("")
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            if ctype == "text/plain":
                try:
                    parts.append(part.get_content())
                except Exception:
                    payload = part.get_payload(decode=True) or b""
                    parts.append(payload.decode("utf-8", errors="replace"))
            elif ctype == "text/html" and len(parts) <= 4:
                try:
                    html = part.get_content()
                except Exception:
                    html = (part.get_payload(decode=True) or b"").decode("utf-8", errors="replace")
                text = re.sub(r"<[^>]+>", " ", str(html))
                parts.append(re.sub(r"\s+", " ", text).strip())
    else:
        try:
            parts.append(str(msg.get_content()))
        except Exception:
            payload = msg.get_payload(decode=True) or b""
            parts.append(payload.decode("utf-8", errors="replace"))
    return "\n".join(parts).strip()


def _extract_msg(path: Path) -> str:
    raw = path.read_bytes()
    try:
        as_utf16 = raw.decode("utf-16-le", errors="ignore")
        chunks = re.findall(r"[\x20-\x7e\n\r\t]{8,}", as_utf16)
        if chunks:
            return "\n".join(chunks[:200])
    except Exception:
        pass
    text = raw.decode("utf-8", errors="ignore")
    chunks = re.findall(r"[\x20-\x7e\n\r\t]{8,}", text)
    return "\n".join(chunks[:200]) if chunks else text[:8000]


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for i, page in enumerate(reader.pages):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(f"--- page {i + 1} ---\n{t}")
    return "\n".join(parts).strip()


def _image_media_type(path: Path) -> str:
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
    }.get(path.suffix.lower(), "image/jpeg")


def _strip_json_fence(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _candidate_models() -> list[str]:
    preferred = os.environ.get("ANTHROPIC_HAIKU_MODEL", "").strip() or HAIKU_MODEL
    out: list[str] = []
    for c in [preferred, *_HAIKU_FALLBACKS]:
        if c and c not in out:
            out.append(c)
    return out


def _haiku_complete(*, system: str, messages: list[dict], max_tokens: int = 4096) -> tuple[str, str]:
    from agents.llm import complete

    last_err: Exception | None = None
    for model in _candidate_models():
        try:
            return complete(model=model, system=system, messages=messages, max_tokens=max_tokens), model
        except Exception as exc:  # noqa: BLE001
            msg = str(exc).lower()
            if "not_found" in msg or "404" in msg:
                last_err = exc
                continue
            raise
    raise RuntimeError(f"No usable Haiku model from {_candidate_models()}: {last_err}")


def _haiku_vision(*, path: Path, system: str, user_text: str) -> tuple[str, str]:
    from agents.llm import get_client

    b64 = base64.standard_b64encode(path.read_bytes()).decode("ascii")
    media = _image_media_type(path)
    client = get_client()
    last_err: Exception | None = None
    for model in _candidate_models():
        try:
            resp = client.messages.create(
                model=model,
                max_tokens=4096,
                system=system,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image",
                                "source": {"type": "base64", "media_type": media, "data": b64},
                            },
                            {"type": "text", "text": user_text},
                        ],
                    }
                ],
            )
            parts: list[str] = []
            for block in resp.content:
                if getattr(block, "type", None) == "text":
                    parts.append(block.text)
                elif hasattr(block, "text"):
                    parts.append(block.text)
            return "\n".join(parts).strip(), model
        except Exception as exc:  # noqa: BLE001
            msg = str(exc).lower()
            if "not_found" in msg or "404" in msg:
                last_err = exc
                continue
            raise
    raise RuntimeError(f"No usable Haiku vision model from {_candidate_models()}: {last_err}")



def _safe_json_loads(raw: str) -> dict[str, Any]:
    """Parse model JSON; tolerate fences and truncated Claude output."""
    cleaned = _strip_json_fence(raw).strip()
    if not cleaned:
        return {
            "lines": [],
            "questionnaire_answers": [],
            "notes": "empty_model_json",
            "confidence": 0.0,
        }

    candidates = [cleaned]
    # If truncated mid-string, close quote then balance braces/brackets
    if cleaned.count('"') % 2 == 1:
        candidates.append(cleaned + '"')
    # Drop trailing incomplete , "key": fragment
    trimmed = re.sub(r',\s*"[^"]*$', '', cleaned)
    if trimmed != cleaned:
        candidates.append(trimmed)
        if trimmed.count('"') % 2 == 1:
            candidates.append(trimmed + '"')

    tried: list[str] = []
    for base in candidates:
        opens = base.count("{") - base.count("}")
        opens_a = base.count("[") - base.count("]")
        balanced = base + ("]" * max(0, opens_a)) + ("}" * max(0, opens))
        for cand in (base, balanced):
            if cand in tried:
                continue
            tried.append(cand)
            try:
                data = json.loads(cand)
            except json.JSONDecodeError:
                continue
            if not isinstance(data, dict):
                continue
            if cand != cleaned:
                note = str(data.get("notes") or "")
                data["notes"] = (note + " | json_repaired_truncated").strip(" |")
                try:
                    data["confidence"] = min(float(data.get("confidence") or 0.4), 0.45)
                except (TypeError, ValueError):
                    data["confidence"] = 0.35
            return data

    # Last resort: do not crash Parse all
    return {
        "lines": [],
        "questionnaire_answers": [],
        "notes": "json_parse_failed_unrecoverable",
        "confidence": 0.0,
        "extraction_notes": cleaned[:500],
    }


def _call_haiku_text(document_text: str, rfx_ctx: dict[str, Any], *, source_name: str) -> dict[str, Any]:
    user_payload = {
        "source_file": source_name,
        "rfx": rfx_ctx,
        "vendor_document_text": document_text[:120000],
    }
    prompt = (
        "Parse this vendor reply into the JSON schema. "
        "Use the RFx line list only as matching context — do not invent prices.\n\n"
        + json.dumps(user_payload, ensure_ascii=False)
    )
    raw, model_used = _haiku_complete(
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": prompt}],
        max_tokens=8192,
    )
    parsed = _safe_json_loads(raw)
    parsed["_model_used"] = model_used
    return parsed


def _call_haiku_image(path: Path, rfx_ctx: dict[str, Any]) -> dict[str, Any]:
    user_text = (
        "Transcribe then parse this vendor reply image into the required JSON schema. "
        "Extract only visible prices/answers. Use RFx lines as matching context only.\n\n"
        + json.dumps({"source_file": path.name, "rfx": rfx_ctx}, ensure_ascii=False)
    )
    raw, model_used = _haiku_vision(path=path, system=SYSTEM_PROMPT, user_text=user_text)
    parsed = _safe_json_loads(raw)
    parsed["_model_used"] = model_used
    return parsed


def _quote_from_llm(
    path: Path,
    vendor_id: str,
    source_format: str,
    parsed: dict[str, Any],
) -> ExtractedQuote:
    model_used = str(parsed.pop("_model_used", HAIKU_MODEL))
    currency = str(parsed.get("currency") or "INR")
    vendor_name = str(parsed.get("vendor_name") or "")
    raw_lines = parsed.get("lines") or []
    lines = [_normalize_line(dict(row), currency) for row in raw_lines if isinstance(row, dict)]
    answers = _normalize_answers(parsed.get("questionnaire_answers") or [])
    notes = str(parsed.get("notes") or "")
    extraction_notes = str(parsed.get("extraction_notes") or "")
    conf = parsed.get("confidence")
    try:
        confidence = float(conf) if conf is not None else (0.7 if lines else 0.35)
    except (TypeError, ValueError):
        confidence = 0.5
    evidence_extra: list[dict[str, Any]] = []
    for item in parsed.get("raw_evidence") or []:
        if isinstance(item, dict):
            evidence_extra.append(
                {
                    "snippet": str(item.get("snippet") or item.get("text") or "")[:500],
                    "location": str(item.get("location") or ""),
                }
            )
    evidence = _meta_evidence(
        path,
        parse_method=f"claude_haiku:{model_used}",
        vendor_name=vendor_name,
        currency=currency,
        extraction_notes=extraction_notes,
        extra=evidence_extra,
    )
    if extraction_notes and extraction_notes not in notes:
        notes = (notes + " | " + extraction_notes).strip(" |")
    return _quote(
        vendor_id=vendor_id,
        source_format=source_format,
        lines=lines,
        questionnaire_answers=answers,
        notes=notes,
        confidence=max(0.0, min(1.0, confidence)),
        raw_evidence=evidence,
    )


def _parse_unstructured(path: Path, vendor_id: str, rfx_ctx: dict[str, Any]) -> ExtractedQuote:
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXTS:
        parsed = _call_haiku_image(path, rfx_ctx)
        return _quote_from_llm(path, vendor_id, suffix.lstrip(".") or "image", parsed)

    if suffix == ".docx":
        text = _extract_docx(path)
        fmt = "docx"
    elif suffix == ".pdf":
        text = _extract_pdf(path)
        fmt = "pdf"
    elif suffix == ".eml":
        text = _extract_eml(path)
        fmt = "eml"
    elif suffix == ".msg":
        text = _extract_msg(path)
        fmt = "msg"
    else:
        text = _extract_txt(path)
        fmt = suffix.lstrip(".") or "txt"

    if not text.strip():
        return _quote(
            vendor_id=vendor_id,
            source_format=fmt,
            lines=[],
            questionnaire_answers=[],
            notes="No readable text extracted from file.",
            confidence=0.0,
            raw_evidence=_meta_evidence(
                path,
                parse_method=f"claude_haiku:{HAIKU_MODEL}",
                extraction_notes="empty extraction",
            ),
        )

    parsed = _call_haiku_text(text, rfx_ctx, source_name=path.name)
    return _quote_from_llm(path, vendor_id, fmt, parsed)


def parse_response(
    path: Union[str, Path],
    vendor_id: str,
    rfx: Any = None,
) -> ExtractedQuote:
    """Parse one vendor reply into ExtractedQuote.

    JSON/CSV are deterministic. Email/Word/PDF/image/text use Claude Haiku.
    Questionnaire answers are shaped for knockout quality gates when RFx is given.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    if not p.is_file():
        raise ValueError(f"Not a file: {p}")

    vid = _guess_vendor_id(p, vendor_id or None)
    suffix = p.suffix.lower()
    rfx_ctx = _rfx_context(rfx)

    if suffix == ".json":
        quote = _parse_json(p, vid)
    elif suffix == ".csv":
        quote = _parse_csv(p, vid)
    else:
        quote = _parse_unstructured(p, vid, rfx_ctx)

    gated = _quality_gate_answers(list(quote.questionnaire_answers or []), rfx_ctx)
    return quote.model_copy(update={"questionnaire_answers": gated})


def parse_vendor_file(
    path: str,
    *,
    vendor_id: str | None = None,
    vendor_hint: str = "",
    rfx: Any = None,
) -> dict:
    """Thin wrapper returning a plain dict with crew helper fields."""
    p = Path(path)
    vid = _guess_vendor_id(p, vendor_id, vendor_hint)
    quote = parse_response(p, vid, rfx)
    data = quote.model_dump()
    meta: dict[str, Any] = {}
    for item in quote.raw_evidence:
        if isinstance(item, dict) and item.get("kind") == "meta":
            meta = item
            break
    data["vendor_name"] = meta.get("vendor_name", "")
    data["parse_method"] = meta.get("parse_method", "")
    data["source_file"] = meta.get("source_file", p.name)
    data["currency"] = meta.get("currency", "INR")
    data["extraction_notes"] = meta.get("extraction_notes", "")
    return data


def parse_vendor_dir(directory: str) -> list[dict]:
    """Parse every non-hidden supported file in directory (sorted)."""
    d = Path(directory)
    if not d.is_dir():
        raise NotADirectoryError(d)
    out: list[dict] = []
    for path in sorted(d.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.suffix.lower() not in SUPPORTED and path.suffix.lower() not in {".md"}:
            continue
        out.append(parse_vendor_file(str(path)))
    return out


class DocumentParserAgent:
    """Object-oriented façade over parse_response."""

    def __init__(self, rfx: Any = None):
        self.rfx = rfx

    def parse(self, path: Union[str, Path], vendor_id: str = "") -> ExtractedQuote:
        return parse_response(path, vendor_id or "", self.rfx)

    def parse_file(self, path: str, vendor_id: str | None = None, vendor_hint: str = "") -> dict:
        return parse_vendor_file(path, vendor_id=vendor_id, vendor_hint=vendor_hint, rfx=self.rfx)

    def parse_dir(self, directory: str) -> list[dict]:
        return parse_vendor_dir(directory)

    def seed_inbox(self, inbox_dir=None, vendor_dir=None):
        return seed_inbox(inbox_dir=inbox_dir, vendor_dir=vendor_dir)

    def list_inbox(self, inbox_dir=None):
        return list_inbox(inbox_dir=inbox_dir)

    def parse_one(self, path, vendor_id: str = "") -> ExtractedQuote:
        return parse_one(path, vendor_id=vendor_id, rfx=self.rfx)

    def parse_all(self, directory=None, *, vendor_ids=None, skip_llm: bool = False):
        return parse_all(directory, rfx=self.rfx, vendor_ids=vendor_ids, skip_llm=skip_llm)


# ---------------------------------------------------------------------------
# Inbox helpers (wizard stub inbound mail)
# ---------------------------------------------------------------------------

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_VENDOR_DIR = REPO_ROOT / "data" / "vendor_responses"
DEFAULT_INBOX_DIR = REPO_ROOT / "data" / "inbox"

# Primary demo samples only (ignore leftover V1_sri_* stubs).
_PRIMARY_INBOX_SEED: list[tuple[str, str, str]] = [
    # (source_name, dest_name, kind)  kind: structured | email
    ("V01_PackForge_response.json", "inbound_V01_PackForge.json", "structured"),
    ("V02_CartonWorks_response.csv", "inbound_V02_CartonWorks.csv", "structured"),
    ("V03_PacificBoard_email.txt", "inbound_V03_PacificBoard.eml", "email"),
    ("V04_QuickCorr_prose.txt", "inbound_V04_QuickCorr.txt", "email"),
    ("V05_NestPack_messy.txt", "inbound_V05_NestPack.txt", "email"),
]

_EMAIL_META = {
    "V01": ("quotes@packforge.in", "PackForge India"),
    "V02": ("rfq@cartonworks.in", "CartonWorks LLP"),
    "V03": ("tenders@pacificboard.com", "PacificBoard Inc"),
    "V04": ("sales@quickcorr.in", "QuickCorr Pvt Ltd"),
    "V05": ("bid@nestpack.in", "NestPack Industries"),
}


def _vendor_id_from_name(name: str) -> str:
    m = re.search(r"(V\d+)", name, re.I)
    return m.group(1).upper() if m else ""


def _wrap_as_email(body: str, *, vendor_id: str, source_name: str) -> str:
    """Ensure From/To/Subject/Date headers for stub inbound mail."""
    if re.match(r"(?i)^from:\s*", body.lstrip()):
        # Already looks like an email (e.g. V03 sample)
        return body if body.endswith("\n") else body + "\n"
    from_addr, display = _EMAIL_META.get(vendor_id, (f"quotes@{vendor_id.lower()}.example", vendor_id))
    headers = (
        f"From: {display} <{from_addr}>\n"
        f"To: procurement@buyer.example\n"
        f"Subject: Re: RFX-CORR-2026-001 — {display} quotation\n"
        f"Date: Mon, 15 Sep 2026 11:00:00 +0530\n"
        f"X-Stub-Source: data/vendor_responses/{source_name}\n"
        f"\n"
    )
    return headers + body


def seed_inbox(
    inbox_dir: str | Path | None = None,
    vendor_dir: str | Path | None = None,
) -> list[Path]:
    """Create/refresh stub inbound files from primary vendor_responses. Idempotent."""
    inbox = Path(inbox_dir) if inbox_dir else DEFAULT_INBOX_DIR
    vendors = Path(vendor_dir) if vendor_dir else DEFAULT_VENDOR_DIR
    inbox.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []
    for src_name, dest_name, kind in _PRIMARY_INBOX_SEED:
        src = vendors / src_name
        if not src.exists():
            # Soft-skip missing samples so seed stays usable mid-refactor
            continue
        dest = inbox / dest_name
        raw = src.read_text(encoding="utf-8", errors="replace")
        vid = _vendor_id_from_name(src_name) or _vendor_id_from_name(dest_name)
        if kind == "structured":
            dest.write_text(raw, encoding="utf-8")
        else:
            dest.write_text(_wrap_as_email(raw, vendor_id=vid, source_name=src_name), encoding="utf-8")
        written.append(dest)
    return written


def list_inbox(inbox_dir: str | Path | None = None) -> list[dict]:
    """List stub inbound files: path, vendor_id, filename, format."""
    inbox = Path(inbox_dir) if inbox_dir else DEFAULT_INBOX_DIR
    if not inbox.is_dir():
        return []
    rows: list[dict] = []
    for path in sorted(inbox.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        if path.suffix.lower() not in SUPPORTED and path.suffix.lower() not in {".md"}:
            continue
        rows.append(
            {
                "path": str(path),
                "vendor_id": _vendor_id_from_name(path.name) or _guess_vendor_id(path),
                "filename": path.name,
                "format": path.suffix.lower().lstrip(".") or "unknown",
            }
        )
    return rows


def parse_one(path: Union[str, Path], vendor_id: str = "", rfx: Any = None) -> ExtractedQuote:
    """Alias of parse_response for wizard/orchestrator call sites."""
    return parse_response(path, vendor_id or "", rfx)


def parse_all(
    directory: str | Path | None = None,
    rfx: Any = None,
    *,
    vendor_ids: list[str] | None = None,
    skip_llm: bool = False,
) -> list[ExtractedQuote]:
    """Parse every supported file in directory (default data/inbox).

    When skip_llm=True, only JSON/CSV (deterministic) files are parsed — useful
    for fast smoke tests without Haiku calls.
    """
    d = Path(directory) if directory else DEFAULT_INBOX_DIR
    if not d.is_dir():
        raise NotADirectoryError(d)

    allow = {v.upper() for v in vendor_ids} if vendor_ids else None
    out: list[ExtractedQuote] = []
    for path in sorted(d.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED and suffix not in {".md"}:
            continue
        if skip_llm and suffix not in JSON_CSV_EXTS:
            continue
        vid = _guess_vendor_id(path)
        if allow is not None and vid.upper() not in allow:
            # also allow V01 vs V1 soft match
            soft = re.sub(r"^V0+", "V", vid.upper())
            allow_soft = {re.sub(r"^V0+", "V", a) for a in allow}
            if vid.upper() not in allow and soft not in allow_soft:
                continue
        try:
            out.append(parse_response(path, vid, rfx))
        except Exception as exc:  # noqa: BLE001 — one bad file must not fail the batch
            out.append(
                _quote(
                    vendor_id=vid or "UNKNOWN",
                    source_format=suffix.lstrip(".") or "unknown",
                    lines=[],
                    questionnaire_answers=[],
                    notes=f"parse_failed: {exc}",
                    confidence=0.0,
                    raw_evidence=_meta_evidence(
                        path,
                        parse_method="parse_all_error",
                        vendor_name="",
                        currency="",
                        extraction_notes=str(exc),
                    ),
                )
            )
    return out


# Alias expected by some __init__/UI imports
DocumentParserAgent = DocumentParserAgent
